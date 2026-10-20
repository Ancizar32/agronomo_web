#!/usr/bin/env python3
"""Genera el SQL de carga inicial de productores, predios y lotes desde el
Excel "Info Agrónomo" entregado por el cliente.

Uso:
    python3 database/tools/importar_info_agronomo.py ARCHIVO.xlsx DIRECTORIO_SALIDA

Produce dos archivos en DIRECTORIO_SALIDA:
    carga_info_agronomo.sql      Script a ejecutar después de la migración 034.
    carga_info_agronomo_reporte.md  Decisiones tomadas y datos que requieren seguimiento.

Reglas de la carga:
- Cada fila del Excel es un lote. Los lotes se agrupan en un predio por
  productor + nombre de predio + vereda.
- Del predio se toma el contrato más reciente, el anticipo sumado y el GPS del
  primer lote que lo tenga. Las diferencias entre lotes quedan en el reporte.
- "Número de plantas /Ha" del Excel es el total de plantas del lote; en la BD
  se guarda por hectárea (total / hectáreas).
- Las proyecciones van a lote_proyecciones con kg = kg/planta × plantas actuales.
- Los IDs se derivan del contenido (md5), así que el script es repetible: volver
  a ejecutarlo actualiza los mismos registros en lugar de duplicarlos.
- Departamento, municipio, vereda, técnicos y cultivo se resuelven por nombre
  dentro de MySQL, para que el mismo SQL sirva en local y en producción.
"""
import datetime
import hashlib
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

CULTIVO = 'Gulupa'
ACTOR = 'carga_excel'
TECNICO_NO_APLICA = re.compile(r'^no aplica', re.I)
DMS = re.compile(r'\s*(\d+)\s*°\s*(\d+)\s*\'\s*([\d.]+)\s*"\s*([NSEW])\s*', re.I)


def uid(*parts):
    return hashlib.md5('|'.join(str(p) for p in parts).encode()).hexdigest()


def clean(v):
    if v is None:
        return None
    if isinstance(v, str):
        v = re.sub(r'\s+', ' ', v).strip()
        return v or None
    return v


def yes(v):
    return str(v or '').strip().lower() in ('si', 'sí')


def to_date(v):
    return v.date() if isinstance(v, datetime.datetime) else None


def dms_to_decimal(v):
    m = DMS.fullmatch(str(v or ''))
    if not m:
        return None
    d = int(m[1]) + int(m[2]) / 60 + float(m[3]) / 3600
    return round(-d if m[4].upper() in 'SW' else d, 7)


def nit_dv(nit):
    pesos = [3, 7, 13, 17, 19, 23, 29, 37, 41, 43, 47, 53, 59, 67, 71]
    t = sum(int(c) * pesos[i] for i, c in enumerate(str(nit)[::-1]))
    r = t % 11
    return r if r < 2 else 11 - r


def sql(v):
    if v is None:
        return 'NULL'
    if isinstance(v, bool):
        return '1' if v else '0'
    if isinstance(v, (int, float)):
        return repr(v)
    if isinstance(v, datetime.date):
        return f"'{v.isoformat()}'"
    return "'" + str(v).replace('\\', '\\\\').replace("'", "''") + "'"


def leer(path):
    wb = openpyxl.load_workbook(path)
    ws = wb['Soft. Agronomos'] if 'Soft. Agronomos' in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    headers = [clean(h) or '' for h in rows[0]]
    out = []
    for i, r in enumerate(rows[1:], start=2):
        if any(v is not None for v in r):
            out.append((i, {h: clean(v) for h, v in zip(headers, r) if h}))
    return out


def construir(filas):
    """Agrupa las filas en productores y predios. Lo usan este script y crear_tecnicos_info_agronomo.py."""
    reporte = defaultdict(list)

    # ---------------- Productores (clave: nombre normalizado) ----------------
    productores = {}
    for n, r in filas:
        nombre = r['Nombre Productor']
        key = nombre.upper()
        p = productores.setdefault(key, {'id': uid('productor', key), 'nombre': nombre, 'tipo': r['Tipo productor'],
                                         'cedulas': Counter(), 'nit': None, 'dv': None,
                                         'telefono': None, 'correos': [], 'filas': []})
        p['filas'].append(n)
        if r.get('Cédula Ciudadanía'):
            p['cedulas'][str(r['Cédula Ciudadanía'])] += 1
        if r.get('Nit'):
            p['nit'] = str(r['Nit'])
            p['dv'] = str(r['DV']) if r.get('DV') is not None else None
        if r.get('Número Contacto') and not p['telefono']:
            p['telefono'] = str(r['Número Contacto'])
        for c in re.split(r'[,;\s]+', r.get('Correo') or ''):
            if c and c not in p['correos']:
                p['correos'].append(c)

    for p in productores.values():
        p['cedula'] = p['cedulas'].most_common(1)[0][0] if p['cedulas'] else None
        if len(p['cedulas']) > 1:
            reporte['Productores con varias cédulas'].append(
                f"{p['nombre']}: se cargó {p['cedula']}; también aparece {', '.join(c for c in p['cedulas'] if c != p['cedula'])} "
                f"(filas {', '.join(map(str, p['filas']))}). Pendiente que el cliente confirme.")
        if p['nit'] and p['dv'] is None:
            p['dv'] = str(nit_dv(p['nit']))
            reporte['DV calculado (algoritmo DIAN)'].append(f"{p['nombre']}: NIT {p['nit']} → DV {p['dv']}")
        if len(p['correos']) > 1:
            reporte['Correos adicionales no cargados'].append(
                f"{p['nombre']}: se cargó {p['correos'][0]}; adicionales: {', '.join(p['correos'][1:])}")

    # ---------------- Predios (productor + predio + vereda) ----------------
    predios = {}
    for n, r in filas:
        nombre_predio = re.sub(r'\s+L\d+$', '', r['Nombre Predio'], flags=re.I)
        prod = productores[r['Nombre Productor'].upper()]
        key = (prod['id'], nombre_predio.upper(), (r.get('Vereda') or '').upper())
        f = predios.setdefault(key, {'id': uid('finca', *key), 'productor': prod, 'nombre': nombre_predio,
                                     'departamento': r['Departamento'], 'ciudad': r['Ciudad'],
                                     'vereda': r.get('Vereda'), 'lotes': []})
        f['lotes'].append((n, r))

    for f in predios.values():
        lotes = f['lotes']
        etiqueta = f"{f['nombre']} ({f['productor']['nombre']})"

        con_fecha = [(to_date(r.get('Fecha de Contrato')), r) for _, r in lotes if to_date(r.get('Fecha de Contrato'))]
        contrato = max(con_fecha, key=lambda x: x[0]) if con_fecha else (None, {})
        f['contrato'] = any(yes(r.get('Contrato de Proveeduría')) for _, r in lotes)
        f['fecha_contrato'] = contrato[0]
        f['version_contrato'] = contrato[1].get('Versión de Contrato') or next(
            (r.get('Versión de Contrato') for _, r in lotes if r.get('Versión de Contrato')), None)
        if len({d for d, _ in con_fecha}) > 1:
            reporte['Predios con contratos distintos por lote (se tomó el más reciente)'].append(
                f"{etiqueta}: " + '; '.join(f"{r['Lote']} {to_date(r.get('Fecha de Contrato')) or 'sin fecha'}"
                                             f" {r.get('Versión de Contrato') or ''}".strip() for _, r in lotes))

        ica = next((r for _, r in lotes if r.get('No. ICA')), {})
        f['registro_ica'] = any(yes(r.get('Reg. ICA')) for _, r in lotes)
        f['numero_ica'] = str(ica['No. ICA']) if ica.get('No. ICA') else None
        f['vencimiento_ica'] = to_date(ica.get('Fecha vencimiento ICA'))

        valores = [r['Valor Anticipo'] for _, r in lotes if r.get('Valor Anticipo')]
        f['anticipo'] = any(str(r.get('Anticipo')).upper() == 'SI' for _, r in lotes)
        f['valor_anticipo'] = sum(valores) if valores else None
        if len(valores) > 1:
            reporte['Anticipos sumados de varios lotes'].append(f"{etiqueta}: {' + '.join(map(str, valores))}")
        if f['anticipo'] and not valores:
            reporte['Anticipo = SI sin valor'].append(etiqueta)

        gps = next((r for _, r in lotes if r.get('GPS-latitud') and str(r.get('Estatus')).lower() == 'activo'), None) \
            or next((r for _, r in lotes if r.get('GPS-latitud')), {})
        f['latitud'] = dms_to_decimal(gps.get('GPS-latitud'))
        f['longitud'] = dms_to_decimal(gps.get('GPS-longitud'))
        f['url'] = gps.get('URL Localización')
        if len({(r.get('GPS-latitud'), r.get('GPS-longitud')) for _, r in lotes}) > 1:
            reporte['Predios con GPS distinto por lote (se tomó el del primer lote activo)'].append(
                f"{etiqueta}: {gps.get('Lote')} {gps.get('GPS-latitud')} {gps.get('GPS-longitud')}")

        f['hectareas'] = round(sum(r.get('Has. Totales') or 0 for _, r in lotes), 4)
        f['estado'] = 'ACTIVO' if any(str(r.get('Estatus')).lower() == 'activo' for _, r in lotes) else 'INACTIVO'

        tecnicos = Counter(r['Asistente Técnico'] for _, r in lotes
                           if r.get('Asistente Técnico') and not TECNICO_NO_APLICA.match(r['Asistente Técnico']))
        f['tecnicos'] = [t for t, _ in tecnicos.most_common()]
        if len(tecnicos) > 1:
            reporte['Predios con varios técnicos (se asignan todos)'].append(f"{etiqueta}: {', '.join(f['tecnicos'])}")
        if not tecnicos:
            reporte['Predios sin técnico asignado'].append(
                f"{etiqueta}: {', '.join(sorted({r.get('Asistente Técnico') or 'vacío' for _, r in lotes}))}")

        certs = {}
        for codigo, col, vence in (('GLOBALGAP', 'GlobalG.A.P.', 'Valido Hasta 1'),
                                   ('RAINFOREST', 'RainForest', 'Valido Hasta 2'),
                                   ('FAIRTRADE', 'Fairtrade', 'Valido Hasta 3')):
            fechas = [to_date(r.get(vence)) for _, r in lotes if yes(r.get(col))]
            if fechas:
                certs[codigo] = max((d for d in fechas if d), default=None)
        micro = [to_date(r.get('Vencimiento análisis Microbiológico')) for _, r in lotes
                 if to_date(r.get('Vencimiento análisis Microbiológico'))]
        if micro:
            certs['MICROBIOLOGICO'] = max(micro)
        f['certs'] = certs

    return productores, predios, reporte


def main(xlsx, outdir):
    filas = leer(xlsx)
    productores, predios, reporte = construir(filas)

    # ---------------- SQL ----------------
    sentencias = []
    w = sentencias.append
    w('-- Carga inicial generada por database/tools/importar_info_agronomo.py')
    w(f'-- Origen: {Path(xlsx).name} · {len(filas)} lotes · {len(predios)} predios · {len(productores)} productores')
    w('-- Requiere la migración 034_catalogos_carga_predios.sql. Es repetible.')
    w('SET NAMES utf8mb4;')
    w('START TRANSACTION;')
    w('')
    w('-- Cultivo de todos los lotes')
    w(f"INSERT INTO cultivos (id, descripcion, codigo, sync, voided, created_at, updated_at, created_by)\n"
      f"SELECT {sql(uid('cultivo', CULTIVO))}, {sql(CULTIVO)}, NULL, '1', '1', NOW(), NOW(), {sql(ACTOR)}\n"
      f"FROM DUAL WHERE NOT EXISTS (SELECT 1 FROM cultivos WHERE descripcion = {sql(CULTIVO)});")
    w(f"SET @cultivo = (SELECT id FROM cultivos WHERE descripcion = {sql(CULTIVO)} ORDER BY voided DESC LIMIT 1);")
    w('')
    w('-- Productores')
    for p in productores.values():
        w(f"INSERT INTO productores (id, tipo, nombre, cedula, nit, dv, telefono, correo, activo) VALUES ("
          f"{sql(p['id'])}, {sql(p['tipo'])}, {sql(p['nombre'])}, {sql(p['cedula'])}, {sql(p['nit'])}, {sql(p['dv'])}, "
          f"{sql(p['telefono'])}, {sql(p['correos'][0] if p['correos'] else None)}, 1)\n"
          f"  ON DUPLICATE KEY UPDATE tipo=VALUES(tipo), nombre=VALUES(nombre), cedula=VALUES(cedula), nit=VALUES(nit), "
          f"dv=VALUES(dv), telefono=VALUES(telefono), correo=VALUES(correo);")

    for f in predios.values():
        w('')
        w(f"-- Predio {f['nombre']} · {f['productor']['nombre']} · {len(f['lotes'])} lote(s)")
        w(f"SET @dep = (SELECT id FROM geo_departamentos WHERE nombre = {sql(f['departamento'])} LIMIT 1);")
        w(f"SET @mun = (SELECT id FROM geo_municipios WHERE departamento_id = @dep AND nombre = {sql(f['ciudad'])} LIMIT 1);")
        # El Excel escribe algunas veredas sin espacios o sin artículo
        # ("QUEBRADABONITA", "CHANCOS") frente al catálogo DANE.
        v = sql(f['vereda'])
        w(f"SET @loc = (SELECT id FROM geo_localidades_rurales WHERE municipio_id = @mun AND activo = 1 AND ("
          f"nombre = {v} OR REPLACE(nombre, ' ', '') = REPLACE({v}, ' ', '') OR nombre IN "
          f"(CONCAT('LA ', {v}), CONCAT('EL ', {v}), CONCAT('LAS ', {v}), CONCAT('LOS ', {v})))\n"
          f"  ORDER BY nombre = {v} DESC, tipo = 'VEREDA' DESC LIMIT 1);")
        w("SET @dep_nombre = (SELECT nombre FROM geo_departamentos WHERE id = @dep);")
        w("SET @mun_nombre = (SELECT nombre FROM geo_municipios WHERE id = @mun);")
        w(f"SET @vereda = COALESCE((SELECT nombre FROM geo_localidades_rurales WHERE id = @loc), {sql(f['vereda'])});")
        tecnico = f['tecnicos'][0] if f['tecnicos'] else None
        w(f"SET @tecnico = (SELECT id FROM `user` WHERE TRIM(name) = {sql(tecnico)} AND void = '1' LIMIT 1);"
          if tecnico else "SET @tecnico = NULL;")
        w(f"INSERT INTO fincas (id, productor_id, descripcion, ubicacion, tecnico_id, departamento, ciudad, vereda, "
          f"departamento_id, municipio_id, localidad_rural_id, estado_predio, hectareas_totales, latitud, longitud, "
          f"url_localizacion, contrato_proveeduria, fecha_contrato, version_contrato, registro_ica, numero_ica, "
          f"vencimiento_ica, anticipo, valor_anticipo, sync, voided, created_at, updated_at, created_by) VALUES (\n"
          f"  {sql(f['id'])}, {sql(f['productor']['id'])}, {sql(f['nombre'])}, "
          f"CONCAT_WS(', ', @vereda, COALESCE(@mun_nombre, {sql(f['ciudad'])}), COALESCE(@dep_nombre, {sql(f['departamento'])})), "
          f"@tecnico, COALESCE(@dep_nombre, {sql(f['departamento'])}), COALESCE(@mun_nombre, {sql(f['ciudad'])}), @vereda,\n"
          f"  @dep, @mun, @loc, {sql(f['estado'])}, {sql(f['hectareas'])}, {sql(f['latitud'])}, {sql(f['longitud'])}, "
          f"{sql(f['url'])}, {sql(f['contrato'])}, {sql(f['fecha_contrato'])}, {sql(f['version_contrato'])}, "
          f"{sql(f['registro_ica'])}, {sql(f['numero_ica'])}, {sql(f['vencimiento_ica'])}, {sql(f['anticipo'])}, "
          f"{sql(f['valor_anticipo'])}, '1', '1', NOW(), NOW(), {sql(ACTOR)})\n"
          f"  ON DUPLICATE KEY UPDATE productor_id=VALUES(productor_id), descripcion=VALUES(descripcion), "
          f"ubicacion=VALUES(ubicacion), tecnico_id=VALUES(tecnico_id), departamento=VALUES(departamento), "
          f"ciudad=VALUES(ciudad), vereda=VALUES(vereda), departamento_id=VALUES(departamento_id), "
          f"municipio_id=VALUES(municipio_id), localidad_rural_id=VALUES(localidad_rural_id), "
          f"estado_predio=VALUES(estado_predio), hectareas_totales=VALUES(hectareas_totales), latitud=VALUES(latitud), "
          f"longitud=VALUES(longitud), url_localizacion=VALUES(url_localizacion), "
          f"contrato_proveeduria=VALUES(contrato_proveeduria), fecha_contrato=VALUES(fecha_contrato), "
          f"version_contrato=VALUES(version_contrato), registro_ica=VALUES(registro_ica), numero_ica=VALUES(numero_ica), "
          f"vencimiento_ica=VALUES(vencimiento_ica), anticipo=VALUES(anticipo), valor_anticipo=VALUES(valor_anticipo), "
          f"sync='1', updated_at=NOW();")
        for t in f['tecnicos']:
            w(f"INSERT IGNORE INTO usuario_fincas (usuario_id, finca_id, created_by) "
              f"SELECT id, {sql(f['id'])}, {sql(ACTOR)} FROM `user` WHERE TRIM(name) = {sql(t)} AND void = '1';")
        for codigo, vence in f['certs'].items():
            w(f"INSERT INTO predio_certificaciones (id, finca_id, tipo, vigente, valido_hasta) VALUES ("
              f"{sql(uid('cert', f['id'], codigo))}, {sql(f['id'])}, {sql(codigo)}, 1, {sql(vence)})\n"
              f"  ON DUPLICATE KEY UPDATE vigente=1, valido_hasta=VALUES(valido_hasta);")

        for n, r in f['lotes']:
            lote_id = uid('lote', f['id'], r['Lote'].upper())
            ha = r.get('Has. Totales') or 0
            activo = str(r.get('Estatus')).lower() == 'activo'
            sembradas, actuales = r.get('Número de plantas /Ha sembradas'), r.get('Número de plantas /Ha actuales')
            por_ha = lambda total: round(total / ha) if total and ha else None
            estado = r.get('Estado') if activo else None
            w(f"INSERT INTO lotes (id, finca_id, cultivo_id, nombre, hectareas, area_m2, distancia_surcos, "
              f"distancia_plantas, plantas_ha_sembradas, fecha_siembra, plantas_ha_actuales, fecha_conteo, "
              f"estado_fenologico, origen_material, sync, voided, created_at, updated_at, created_by) VALUES (\n"
              f"  {sql(lote_id)}, {sql(f['id'])}, @cultivo, {sql(r['Lote'])}, {sql(round(ha, 2) if ha else None)}, "
              f"{sql(round(ha * 10000, 2) if ha else None)}, {sql(r.get('Distancia entre Surco'))}, "
              f"{sql(r.get('Distancia entre Plantas'))}, {sql(por_ha(sembradas))}, {sql(to_date(r.get('Fecha de siembra')))}, "
              f"{sql(por_ha(actuales))}, {sql(to_date(r.get('Fecha Conteo')))}, {sql(estado)}, "
              f"{sql(r.get('Origen Material Vegetal'))}, '1', {sql('1' if activo else '0')}, NOW(), NOW(), {sql(ACTOR)})\n"
              f"  ON DUPLICATE KEY UPDATE cultivo_id=VALUES(cultivo_id), nombre=VALUES(nombre), hectareas=VALUES(hectareas), "
              f"area_m2=VALUES(area_m2), distancia_surcos=VALUES(distancia_surcos), distancia_plantas=VALUES(distancia_plantas), "
              f"plantas_ha_sembradas=VALUES(plantas_ha_sembradas), fecha_siembra=VALUES(fecha_siembra), "
              f"plantas_ha_actuales=VALUES(plantas_ha_actuales), fecha_conteo=VALUES(fecha_conteo), "
              f"estado_fenologico=VALUES(estado_fenologico), origen_material=VALUES(origen_material), "
              f"voided=VALUES(voided), sync='1', updated_at=NOW();")
            if activo and not ha and (sembradas or actuales):
                reporte['Lotes activos sin hectáreas (no se pudo calcular plantas/ha)'].append(
                    f"Fila {n} · {f['nombre']} {r['Lote']}: {sembradas or actuales} plantas")
            elif activo and (not ha or not r.get('Fecha de siembra') or not r.get('Estado')):
                faltan = [c for c, v in (('hectáreas', ha), ('fecha de siembra', r.get('Fecha de siembra')),
                                         ('estado', r.get('Estado'))) if not v]
                reporte['Lotes activos con datos faltantes'].append(f"Fila {n} · {f['nombre']} {r['Lote']}: {', '.join(faltan)}")
            if sembradas and actuales and actuales > sembradas:
                reporte['Plantas actuales mayores que sembradas'].append(
                    f"Fila {n} · {f['nombre']} {r['Lote']}: {actuales} actuales / {sembradas} sembradas")

            plantas = actuales or sembradas
            numero = r.get('Numero de Cosechas Proyectadas') or 0
            for k in range(1, 6):
                kg = r.get(f'Productividad Proyectada por Cosecha {k} (Kg/pl)')
                if kg is None or k > numero:
                    continue
                w(f"INSERT INTO lote_proyecciones (id, lote_id, numero_cosecha, productividad_kg_planta, kg_proyectados) "
                  f"VALUES ({sql(uid('proy', lote_id, k))}, {sql(lote_id)}, {k}, {sql(kg)}, "
                  f"{sql(round(kg * plantas, 2) if plantas else None)})\n"
                  f"  ON DUPLICATE KEY UPDATE productividad_kg_planta=VALUES(productividad_kg_planta), "
                  f"kg_proyectados=VALUES(kg_proyectados);")

    w('')
    w('COMMIT;')
    w('')
    w('-- Verificación: técnicos del Excel que no existen como usuario activo')
    tecnicos = sorted({t for f in predios.values() for t in f['tecnicos']})
    union = '\n  UNION ALL '.join(f"SELECT {sql(t)} AS tecnico" for t in tecnicos)
    w(f"SELECT x.tecnico AS tecnico_sin_usuario FROM (\n  {union}\n) x\n"
      f"WHERE NOT EXISTS (SELECT 1 FROM `user` u WHERE TRIM(u.name) = x.tecnico AND u.void = '1');")
    w('-- Verificación: predios cuya ubicación no se encontró en el catálogo DANE')
    w(f"SELECT descripcion, departamento, ciudad, vereda, departamento_id, municipio_id, localidad_rural_id "
      f"FROM fincas WHERE created_by = {sql(ACTOR)} AND (municipio_id IS NULL OR localidad_rural_id IS NULL) ORDER BY ciudad, vereda;")
    w(f"SELECT 'productores' AS tabla, COUNT(*) AS total FROM productores WHERE id IN ("
      + ', '.join(sql(p['id']) for p in productores.values()) + ")\n"
      f"UNION ALL SELECT 'fincas', COUNT(*) FROM fincas WHERE created_by = {sql(ACTOR)}\n"
      f"UNION ALL SELECT 'lotes', COUNT(*) FROM lotes WHERE created_by = {sql(ACTOR)}\n"
      f"UNION ALL SELECT 'lotes activos', COUNT(*) FROM lotes WHERE created_by = {sql(ACTOR)} AND voided = '1';")

    out = Path(outdir)
    out.mkdir(parents=True, exist_ok=True)
    (out / 'carga_info_agronomo.sql').write_text('\n'.join(sentencias) + '\n', encoding='utf-8')

    activos = sum(1 for _, r in filas if str(r.get('Estatus')).lower() == 'activo')
    md = [f'# Reporte de carga · {Path(xlsx).name}', '',
          f'- Lotes: {len(filas)} ({activos} activos, {len(filas) - activos} erradicados, cargados como inactivos)',
          f'- Predios: {len(predios)}', f'- Productores: {len(productores)}',
          f'- Técnicos referenciados: {", ".join(tecnicos)}', f'- Cultivo asignado: {CULTIVO}', '']
    for titulo, items in reporte.items():
        md += [f'## {titulo} ({len(items)})', ''] + [f'- {i}' for i in items] + ['']
    (out / 'carga_info_agronomo_reporte.md').write_text('\n'.join(md), encoding='utf-8')
    print(f'{len(filas)} lotes, {len(predios)} predios, {len(productores)} productores → {out}')


if __name__ == '__main__':
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
